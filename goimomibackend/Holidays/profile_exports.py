"""Staff-only profile exports, shared by the Global Horizons API."""
from io import BytesIO
from xml.sax.saxutils import escape

from django.http import HttpResponse
from django.utils import timezone


FIELDS = (
    ('full_name', 'Full Name'), ('city', 'City'), ('country', 'Country'),
    ('profession', 'Business / Profession'), ('organization', 'Organization / Brand Name'),
    ('years_of_experience', 'Years of experience'), ('interests', 'Areas of interest or expertise'),
    ('website', 'Website'), ('email', 'Email address'), ('connections_sought', 'Connections sought in Sri Lanka'),
    ('ticket_status', 'Ticket booking status'), ('arrival_date', 'Arrival Date'),
    ('arrival_flight_no', 'Arrival Flight No'), ('arrival_time', 'Arrival Time'),
    ('return_date', 'Return Date'), ('return_flight_no', 'Return Flight No'), ('return_time', 'Return Time'),
    ('photo', 'Photo URL'), ('created_at', 'Submitted (UTC)'),
)


def profile_values(profile, request):
    values = []
    for field, _ in FIELDS:
        value = getattr(profile, field)
        if field == 'photo':
            value = request.build_absolute_uri(value.url) if value else ''
        elif field == 'ticket_status':
            value = profile.get_ticket_status_display()
        elif field == 'created_at':
            value = value.isoformat()
        values.append('' if value is None else value)
    return values


def export_profiles(profiles, request, file_type):
    buffer = BytesIO()
    if file_type == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Participant Profiles'
        sheet.append([label for _, label in FIELDS])
        for profile in profiles:
            sheet.append(profile_values(profile, request))
            for cell in sheet[sheet.max_row]:
                # User-supplied strings must remain text, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='14532D')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 32
        for index, (field, _) in enumerate(FIELDS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 45 if field in ('interests', 'connections_sought', 'photo') else 25
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(buffer)
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

        styles = getSampleStyleSheet()
        styles['BodyText'].fontSize = 9
        styles['BodyText'].leading = 13
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        story = [Paragraph('Global Horizons - Srilanka', styles['Title']),
                 Paragraph(f'Participant profiles: {len(profiles)} | Exported {timezone.now():%Y-%m-%d %H:%M} UTC', styles['BodyText']), Spacer(1, 16)]
        for index, profile in enumerate(profiles):
            if index:
                story.append(PageBreak())
            story.append(Paragraph(escape(profile.full_name), styles['Heading2']))
            rows = [[Paragraph(escape(label), styles['BodyText']),
                     Paragraph(escape(str(value)).replace('\n', '<br/>') or '-', styles['BodyText'])]
                    for (_, label), value in zip(FIELDS, profile_values(profile, request))]
            table = Table(rows, colWidths=[150, doc.width - 150], splitInRow=1)
            table.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#edf7f0')),
                ('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.HexColor('#dce4df')),
                ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            story.append(table)

        def footer(canvas, document):
            canvas.saveState()
            canvas.setFont('Helvetica', 8)
            canvas.drawRightString(A4[0] - 36, 20, f'Page {document.page}')
            canvas.restoreState()

        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        content_type = 'application/pdf'
    response = HttpResponse(buffer.getvalue(), content_type=content_type)
    response['Content-Disposition'] = f'attachment; filename="global-horizons-srilanka-{timezone.now():%Y-%m-%d}.{file_type}"'
    response['Cache-Control'] = 'private, no-store'
    return response
