from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib import admin
from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from PIL import Image
from rest_framework.test import APIClient

from .models import GlobalHorizonsProfile


def photo_file(format='PNG'):
    image = BytesIO()
    Image.new('RGB', (12, 12), 'green').save(image, format=format)
    return SimpleUploadedFile(f'portrait.{format.lower()}', image.getvalue(), content_type=f'image/{format.lower()}')


class GlobalHorizonsProfileTests(TestCase):
    def test_staff_exports_filter_profiles_and_produce_real_files(self):
        from openpyxl import load_workbook
        self.client.post(self.url, self.payload(full_name='=1+1', interests='Textiles & design <partners>'), format='multipart')
        self.client.post(self.url, self.payload(full_name='Other Participant', interests='Agriculture'), format='multipart')
        url = reverse('global-horizons-srilanka-export')
        for user in [None, self.customer]:
            self.client.force_authenticate(user=user)
            for file_type in ['xlsx', 'pdf']:
                self.assertIn(self.client.get(url, {'file_type': file_type}).status_code, (401, 403))
        self.client.force_authenticate(user=self.staff)
        response = self.client.get(url, {'file_type': 'xlsx', 'search': 'textiles'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, '=1+1')
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertIn('Arrival Flight No', [cell.value for cell in sheet[1]])
        response = self.client.get(url, {'file_type': 'pdf', 'search': 'textiles'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assertEqual(self.client.get(url, {'file_type': 'csv'}).status_code, 400)
        GlobalHorizonsProfile.objects.filter(full_name='=1+1').update(interests='Long profile description. ' * 200)
        self.assertEqual(self.client.get(url, {'file_type': 'pdf'}).status_code, 200)

    def setUp(self):
        # Each test gets a fresh throttle counter, independent of other API tests.
        cache.clear()
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        settings_override = override_settings(MEDIA_ROOT=directory.name)
        settings_override.enable()
        self.addCleanup(settings_override.disable)
        self.client = APIClient()
        self.url = reverse('global-horizons-srilanka-list')
        self.staff = User.objects.create_user('organizer', is_staff=True)
        self.customer = User.objects.create_user('participant')

    def payload(self, **overrides):
        return {
            'full_name': 'Tamil Participant', 'city': 'Chennai', 'country': 'India',
            'profession': 'Entrepreneur', 'photo': photo_file(), 'organization': 'Example Brand',
            'years_of_experience': 12, 'interests': 'Sustainable textiles',
            'website': 'https://example.com', 'email': 'participant@example.com',
            'connections_sought': 'Meet distributors and potential partners in Colombo.',
            'ticket_status': 'not_booked',
            **overrides,
        }

    def test_public_submission_persists_every_field_and_photo(self):
        payload = self.payload()
        response = self.client.post(self.url, payload, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        profile = GlobalHorizonsProfile.objects.get()
        for key, value in payload.items():
            if key != 'photo':
                self.assertEqual(getattr(profile, key), value)
        self.assertTrue(profile.photo.storage.exists(profile.photo.name))
        self.assertTrue(profile.photo.name.startswith('global_horizons/srilanka/'))
        with profile.photo.open('rb') as saved:
            self.assertEqual(Image.open(saved).size, (12, 12))

    def test_bare_social_link_is_saved_as_https(self):
        response = self.client.post(self.url, self.payload(website='  linkedin.com/in/example?ref=profile#about  '), format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(GlobalHorizonsProfile.objects.get().website, 'https://linkedin.com/in/example?ref=profile#about')

    def test_optional_website_and_zero_years_are_accepted(self):
        payload = self.payload(years_of_experience=0)
        del payload['website']
        response = self.client.post(self.url, payload, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)

    def test_organization_and_ticket_status_are_required(self):
        for field in ['organization', 'ticket_status']:
            for value in [None, '', '   ']:
                with self.subTest(field=field, value=value):
                    payload = self.payload()
                    if value is None:
                        del payload[field]
                    else:
                        payload[field] = value
                    response = self.client.post(self.url, payload, format='multipart')
                    self.assertEqual(response.status_code, 400, response.data)
                    self.assertIn(field, response.data)

    def test_required_fields_email_url_and_experience_validation(self):
        for key, value in [('full_name', '  '), ('city', ''), ('country', ''), ('profession', ''),
                           ('interests', ''), ('connections_sought', ''), ('email', 'invalid'),
                           ('website', 'javascript:alert(1)'), ('years_of_experience', -1),
                           ('years_of_experience', 101), ('years_of_experience', '1.5')]:
            with self.subTest(key=key, value=value):
                response = self.client.post(self.url, self.payload(**{key: value}), format='multipart')
                self.assertEqual(response.status_code, 400, response.data)
                self.assertIn(key, response.data)
        payload = self.payload()
        del payload['photo']
        self.assertEqual(self.client.post(self.url, payload, format='multipart').status_code, 400)
        self.assertFalse(GlobalHorizonsProfile.objects.exists())

    def test_invalid_unsupported_and_oversized_photos_are_rejected(self):
        large = photo_file()
        large = SimpleUploadedFile('large.png', large.read() + b'0' * (5 * 1024 * 1024), content_type='image/png')
        for photo in [SimpleUploadedFile('fake.png', b'not an image', content_type='image/png'), photo_file('GIF'), large]:
            with self.subTest(filename=photo.name):
                response = self.client.post(self.url, self.payload(photo=photo), format='multipart')
                self.assertEqual(response.status_code, 400, response.data)
                self.assertIn('photo', response.data)
        self.assertFalse(GlobalHorizonsProfile.objects.exists())

    def test_only_staff_can_list_read_edit_and_delete_profiles(self):
        created = self.client.post(self.url, self.payload(), format='multipart')
        detail = reverse('global-horizons-srilanka-detail', args=[created.data['id']])
        for user in [None, self.customer]:
            self.client.force_authenticate(user=user)
            for method, url, data in [('get', self.url, {}), ('get', detail, {}),
                                       ('patch', detail, {'city': 'Colombo'}), ('delete', detail, {})]:
                with self.subTest(user=user, method=method):
                    self.assertIn(getattr(self.client, method)(url, data).status_code, (401, 403))
        self.client.force_authenticate(user=self.staff)
        self.assertEqual(len(self.client.get(self.url).data), 1)
        self.assertEqual(self.client.get(detail).status_code, 200)
        old_photo = GlobalHorizonsProfile.objects.get().photo.name
        response = self.client.patch(detail, {'city': 'Colombo'}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(GlobalHorizonsProfile.objects.get().photo.name, old_photo)
        self.assertEqual(GlobalHorizonsProfile.objects.get().city, 'Colombo')
        response = self.client.patch(detail, {'photo': photo_file()}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertNotEqual(GlobalHorizonsProfile.objects.get().photo.name, old_photo)
        self.assertEqual(self.client.delete(detail).status_code, 204)

    def test_django_admin_exposes_requested_section_and_form(self):
        self.assertEqual(GlobalHorizonsProfile._meta.verbose_name_plural, 'Global Horizons - Srilanka')
        self.assertIn(GlobalHorizonsProfile, admin.site._registry)
        self.client.force_login(User.objects.create_superuser('admin', password='test-password'))
        response = self.client.get(reverse('admin:Holidays_globalhorizonsprofile_add'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'connections_sought')

    def test_booked_ticket_requires_and_persists_all_flight_details(self):
        response = self.client.post(self.url, self.payload(ticket_status='booked'), format='multipart')
        self.assertEqual(response.status_code, 400)
        flight_data = dict(arrival_date='2026-10-22', arrival_flight_no='UL 122', arrival_time='12:30',
                           return_date='2026-10-24', return_flight_no='UL 123', return_time='18:45')
        self.assertEqual(set(response.data), set(flight_data))
        response = self.client.post(self.url, self.payload(ticket_status='booked', **flight_data), format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['arrival_flight_no'], 'UL 122')
        self.assertEqual(response.data['return_time'], '18:45:00')
        detail = reverse('global-horizons-srilanka-detail', args=[response.data['id']])
        self.client.force_authenticate(self.staff)
        response = self.client.patch(detail, {'return_date': '2026-10-21'}, format='multipart')
        self.assertEqual(response.status_code, 400)
        response = self.client.patch(detail, {'city': 'Madurai'}, format='multipart')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['arrival_date'], '2026-10-22')
        response = self.client.patch(detail, {'ticket_status': 'not_booked'}, format='multipart')
        self.assertEqual(response.status_code, 200)
        profile = GlobalHorizonsProfile.objects.get(pk=response.data['id'])
        for field in flight_data:
            self.assertFalse(getattr(profile, field))
        response = self.client.patch(detail, {'ticket_status': 'booked'}, format='multipart')
        self.assertEqual(response.status_code, 400)


class ProfileLinkValidationTests(SimpleTestCase):
    def test_supported_web_links(self):
        from .profile_links import normalize_profile_link
        for source, expected in [
            ('', ''), ('  ', ''), ('example.com', 'https://example.com'),
            ('www.example.com/about', 'https://www.example.com/about'),
            ('instagram.com/example/', 'https://instagram.com/example/'),
            ('https://wa.me/919876543210', 'https://wa.me/919876543210'),
            ('https://youtu.be/abc?x=1#part', 'https://youtu.be/abc?x=1#part'),
            ('http://example.com', 'http://example.com'),
            ('//linkedin.com/in/example', 'https://linkedin.com/in/example'),
        ]:
            with self.subTest(source=source):
                self.assertEqual(normalize_profile_link(source), expected)

    def test_invalid_and_non_web_links(self):
        from django.core.exceptions import ValidationError
        from .profile_links import normalize_profile_link
        for value in ['javascript:alert(1)', 'data:text/html,test', 'ftp://example.com/file',
                      'mailto:hello@example.com', 'tel:+123456', 'not a link', '/relative',
                      'https://user:password@example.com', 'https://example.com/' + 'a' * 500]:
            with self.subTest(value=value), self.assertRaises(ValidationError):
                normalize_profile_link(value)
