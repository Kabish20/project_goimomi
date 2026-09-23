from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.core import mail
from django.db import transaction
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from .models import GlobalHorizonsProfile, GlobalHorizonsExportEmail
from .profile_notifications import deliver_profile_export, queue_profile_export


@override_settings(GLOBAL_HORIZONS_EXPORT_RECIPIENTS=['one@example.test', 'two@example.test'], GLOBAL_HORIZONS_EXPORT_CC=['support@example.test'])
class ProfileNotificationTests(TestCase):
    def profile(self, name='First'):
        return GlobalHorizonsProfile.objects.create(
            full_name=name, city='Chennai', country='India', profession='Founder',
            photo='global_horizons/example.png', organization='Example',
            years_of_experience=5, interests='Textiles', email='participant@example.test',
            connections_sought='Business connections')

    def test_create_update_delete_queue_after_commit_and_skip_document_saves(self):
        with patch('Holidays.tasks.send_global_horizons_export_task.apply_async') as publish:
            with self.captureOnCommitCallbacks(execute=True):
                profile = self.profile()
                profile.city = 'Colombo'
                profile.save()
                profile.save(update_fields=['profile_booklet', 'attending_poster', 'attending_poster_image'])
                profile.delete()
                publish.assert_not_called()
            self.assertEqual(publish.call_count, 3)
        self.assertEqual(list(GlobalHorizonsExportEmail.objects.order_by('pk').values_list('event', flat=True)),
                         ['created', 'updated', 'deleted'])

    def test_rollback_does_not_publish_or_keep_an_event(self):
        with patch('Holidays.tasks.send_global_horizons_export_task.apply_async') as publish:
            with self.captureOnCommitCallbacks(execute=True):
                with transaction.atomic():
                    self.profile()
                    transaction.set_rollback(True)
            publish.assert_not_called()
        self.assertFalse(GlobalHorizonsExportEmail.objects.exists())

    def test_full_excel_all_recipients_and_duplicate_job(self):
        self.profile('=1+1')
        self.profile('Second')
        event = GlobalHorizonsExportEmail.objects.order_by('pk').first()
        self.assertTrue(deliver_profile_export(event.pk))
        self.assertTrue(deliver_profile_export(event.pk))
        self.assertEqual(len(mail.outbox), 1)
        message = mail.outbox[0]
        self.assertEqual(message.bcc, ['two@example.test'])
        self.assertEqual(message.cc, ['support@example.test'])
        self.assertEqual(message.to, ['one@example.test'])
        sheet = load_workbook(BytesIO(message.attachments[0][1])).active
        self.assertEqual(sheet.max_row, 3)
        rows = list(sheet.values)
        headers = rows[0]
        self.assertIn('Logo URL', headers)
        self.assertIn('Profile Booklet URL', headers)
        self.assertIn('https://goimomi.com/media/global_horizons/example.png', [row[headers.index('Photo URL')] for row in rows[1:]])
        cell = next(cell for row in sheet for cell in row if cell.value == '=1+1')
        self.assertEqual(cell.data_type, 's')
        event.refresh_from_db()
        self.assertEqual(event.status, 'sent')
        self.assertEqual(event.attempts, 1)

    def test_smtp_failure_is_recorded_and_retried(self):
        event = queue_profile_export('requested')
        with patch('Holidays.profile_notifications.EmailMessage.send', side_effect=ConnectionError), self.assertLogs('Holidays.profile_notifications', level='ERROR'):
            self.assertFalse(deliver_profile_export(event.pk))
        event.refresh_from_db()
        self.assertEqual(event.status, 'pending')
        self.assertEqual(event.last_error, 'ConnectionError')
        self.assertGreater(event.next_attempt_at, timezone.now())
        GlobalHorizonsExportEmail.objects.filter(pk=event.pk).update(next_attempt_at=timezone.now() - timedelta(seconds=1))
        self.assertTrue(deliver_profile_export(event.pk))

    def test_broker_failure_keeps_pending_delivery(self):
        with patch('Holidays.tasks.send_global_horizons_export_task.apply_async', side_effect=ConnectionError), self.assertLogs('Holidays.profile_notifications', level='ERROR'), self.captureOnCommitCallbacks(execute=True):
            self.profile()
        self.assertEqual(GlobalHorizonsExportEmail.objects.get().status, 'pending')

    def test_recovery_command_only_requeues_due_pending_events(self):
        from django.core.management import call_command
        event = queue_profile_export('requested')
        with patch('Holidays.tasks.send_global_horizons_export_task.apply_async') as publish:
            call_command('retry_profile_export_emails')
            publish.assert_called_once_with(args=[event.pk], queue='enquiries', retry=False)

    def test_fifth_failure_is_visible_and_not_automatically_requeued(self):
        event = queue_profile_export('requested')
        GlobalHorizonsExportEmail.objects.filter(pk=event.pk).update(attempts=4)
        with patch('Holidays.profile_notifications.EmailMessage.send', return_value=0), self.assertLogs('Holidays.profile_notifications', level='ERROR'):
            self.assertFalse(deliver_profile_export(event.pk))
        event.refresh_from_db()
        self.assertEqual(event.status, 'failed')
        self.assertEqual(event.attempts, 5)

    @override_settings(GLOBAL_HORIZONS_EXPORT_RECIPIENTS=[])
    def test_no_recipients_does_not_send(self):
        self.profile()
        self.assertFalse(GlobalHorizonsExportEmail.objects.exists())
